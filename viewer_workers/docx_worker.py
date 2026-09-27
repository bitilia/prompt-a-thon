import sys, html as _html

def main():
    if len(sys.argv) < 2:
        sys.exit(1)
    path = sys.argv[1]
    with open(path, 'rb') as f:
        if f.read(4)[:2] != b'PK':
            sys.exit(2)
    try:
        from docx import Document
        doc = Document(path)
        paragraphs = [p.text for p in doc.paragraphs]
        tables_html = []
        for table in doc.tables:
            rows_html = []
            for i, row in enumerate(table.rows):
                tag = "th" if i == 0 else "td"
                cells = "".join(f"<{tag}>{_html.escape(c.text)}</{tag}>" for c in row.cells)
                rows_html.append(f"<tr>{cells}</tr>")
            tables_html.append("<table style='border-collapse:collapse;margin:1em 0;width:100%;'>"+"".join(rows_html)+"</table>")
    except Exception as e:
        print(f"ERROR:{e}", file=sys.stderr); sys.exit(3)
    body_html = "".join(
        f"<p>{_html.escape(p)}</p>" if p.strip() else "<br>"
        for p in paragraphs
    )
    out = (
        "<!doctype html><html><head><meta charset=utf-8><style>"
        "body{font-family:Georgia,serif;font-size:14px;line-height:1.7;margin:0;padding:1rem 1.5rem;background:#fff;color:#111;max-width:800px;}"
        "p{margin:0 0 0.6em;} table{border-collapse:collapse;width:100%;margin:1em 0;}"
        "th,td{border:1px solid #ccc;padding:4px 8px;} th{background:#f1f5f9;}"
        "</style></head><body>" + body_html + "".join(tables_html) + "</body></html>"
    )
    sys.stdout.write(out)

if __name__ == "__main__":
    main()
